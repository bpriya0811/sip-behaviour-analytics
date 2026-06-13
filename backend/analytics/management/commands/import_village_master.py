from __future__ import annotations

from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from analytics.models import GeographicLocation, Question


EXPECTED_HEADERS = {
    "district": "District",
    "taluka": "Taluka",
    "village": "Village",
}


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_key(value) -> str:
    return _clean(value).lower().replace(" ", "").replace("_", "")


def _find_columns(row) -> dict[str, int]:
    lookup = {_header_key(value): index for index, value in enumerate(row)}
    columns: dict[str, int] = {}
    for key, label in EXPECTED_HEADERS.items():
        column = lookup.get(_header_key(label))
        if column is None:
            raise CommandError(
                "Village master must contain District, Taluka, and Village columns."
            )
        columns[key] = column
    return columns


def _sync_geographic_questions() -> None:
    rows = list(
        GeographicLocation.objects.order_by("district", "taluka")
        .values_list("district", "taluka")
        .distinct()
    )
    districts: list[str] = []
    talukas_by_district: dict[str, list[str]] = {}
    for district, taluka in rows:
        if district not in districts:
            districts.append(district)
        talukas_by_district.setdefault(district, [])
        if taluka not in talukas_by_district[district]:
            talukas_by_district[district].append(taluka)

    district_question, _ = Question.objects.update_or_create(
        slug="district",
        defaults={
            "step": 1,
            "section": "Demographic Information",
            "prompt": "District",
            "question_type": "dropdown",
            "required": True,
            "active": True,
            "order": 6,
            "metadata": {},
        },
    )
    district_question.options.all().delete()
    for index, district in enumerate(districts):
        district_question.options.create(label=district, value=district, order=index)

    taluka_question, _ = Question.objects.update_or_create(
        slug="taluka",
        defaults={
            "step": 1,
            "section": "Demographic Information",
            "prompt": "Taluka",
            "question_type": "dropdown",
            "placeholder": "Select district first",
            "required": True,
            "active": True,
            "order": 7,
            "metadata": {
                "dependsOn": "district",
                "dependentPlaceholder": "Select district first",
                "optionsByDistrict": talukas_by_district,
            },
        },
    )
    taluka_question.options.all().delete()
    for index, taluka in enumerate(
        taluka
        for district in districts
        for taluka in talukas_by_district.get(district, [])
    ):
        taluka_question.options.create(label=taluka, value=taluka, order=index)

    Question.objects.update_or_create(
        slug="village",
        defaults={
            "step": 1,
            "section": "Demographic Information",
            "prompt": "Village",
            "question_type": "text",
            "placeholder": "Search and select village",
            "required": True,
            "active": True,
            "order": 8,
            "metadata": {
                "dependsOn": "taluka",
                "dependsOnDistrict": "district",
                "autocomplete": True,
                "source": "geographic_locations",
                "dependentPlaceholder": "Select taluka first",
            },
        },
    )


class Command(BaseCommand):
    help = "Import District, Taluka, and Village rows from Village_Master_Sangli_Kolhapur.xlsx."

    def add_arguments(self, parser):
        parser.add_argument(
            "path",
            nargs="?",
            default=str(settings.BASE_DIR / "data" / "Village_Master_Sangli_Kolhapur.xlsx"),
            help="Path to the Village_Master_Sangli_Kolhapur.xlsx workbook.",
        )

    def handle(self, *args, **options):
        source = Path(options["path"]).expanduser()
        if not source.exists():
            raise CommandError(f"Village master workbook not found: {source}")

        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise CommandError("Village master workbook is empty.") from exc

        columns = _find_columns(header)
        seen: set[tuple[str, str, str]] = set()
        locations: list[GeographicLocation] = []

        for row in rows:
            district = _clean(row[columns["district"]])
            taluka = _clean(row[columns["taluka"]])
            village = _clean(row[columns["village"]])
            if not district or not taluka or not village:
                continue
            key = (district, taluka, village)
            if key in seen:
                continue
            seen.add(key)
            locations.append(
                GeographicLocation(district=district, taluka=taluka, village=village)
            )

        if not locations:
            raise CommandError("No valid District/Taluka/Village rows found in workbook.")

        with transaction.atomic():
            GeographicLocation.objects.all().delete()
            GeographicLocation.objects.bulk_create(locations, batch_size=1000)
            _sync_geographic_questions()

        self.stdout.write(
            self.style.SUCCESS(
                f"Imported {len(locations)} village rows from {source.name}."
            )
        )
